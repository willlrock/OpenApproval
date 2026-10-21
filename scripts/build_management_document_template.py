from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "resources" / "documents"

# standard_business_brief with named TEG A4 overrides based on the supplied PDF.
NAVY = "172A4A"
BLUE = "2B4C7E"
MUTED = "687386"
INK = "182235"
LINE = "C8D2E0"
FILL = "EFF3F8"
INPUT_HIGHLIGHT = "cyan"
PAGE_WIDTH_DXA = 11906
CONTENT_WIDTH_DXA = 9980
TABLE_INDENT_DXA = 110
CELL_MARGIN_DXA = 110


def set_font(run, size=10.5, bold=False, color=INK, italic=False, name="Arial"):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=CELL_MARGIN_DXA, start=CELL_MARGIN_DXA,
                     bottom=CELL_MARGIN_DXA, end=CELL_MARGIN_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(index, len(widths) - 1)])
            set_cell_margins(cell)


def set_table_borders(table, color=LINE, size="5", inside=True):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    names = ["top", "left", "bottom", "right"]
    if inside:
        names += ["insideH", "insideV"]
    for edge in names:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_paragraph_border(paragraph, color=NAVY, size="18"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def add_system_value(cell, label, token, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(label + "\n"), size=7.5, bold=True, color=MUTED)
    set_font(p.add_run(token), size=9.5, bold=True, color=NAVY)


def add_section_heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(f"{number}. {title}"), size=10, bold=True, color=NAVY)


def add_fillable(doc, placeholder, min_lines=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.12)
    p.paragraph_format.right_indent = Cm(0.12)
    p.paragraph_format.keep_together = True
    run = p.add_run(f"[[{placeholder}]]")
    set_font(run, size=10.5, italic=True, color=BLUE)
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), INPUT_HIGHLIGHT)
    run._element.get_or_add_rPr().append(highlight)
    for _ in range(min_lines - 1):
        spacer = p.add_run("\n ")
        set_font(spacer, size=10.5, color=BLUE)
        spacer_highlight = OxmlElement("w:highlight")
        spacer_highlight.set(qn("w:val"), INPUT_HIGHLIGHT)
        spacer._element.get_or_add_rPr().append(spacer_highlight)
    return p


def add_page_field(paragraph):
    set_font(paragraph.add_run("Страница "), size=7.5, color=MUTED)
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (fld_char, instr, separate, text, end):
        run._r.append(node)


def build(kind, filename):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header
    header_table = header.add_table(rows=1, cols=2, width=Cm(17.5))
    set_table_geometry(header_table, [4800, 5180])
    left, right = header_table.rows[0].cells
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("TEG"), size=17, bold=True, color=NAVY)
    set_font(p.add_run(" HOLDING\n"), size=10, bold=True, color=NAVY)
    set_font(p.add_run("СИСТЕМА УПРАВЛЕНИЯ ДОКУМЕНТАМИ"), size=6.8, bold=True, color=MUTED)
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("{{DOCUMENT_KIND}}"), size=10.2, bold=True, color=NAVY)
    rule = header.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(0)
    set_paragraph_border(rule)

    footer = section.footer
    footer_table = footer.add_table(rows=1, cols=2, width=Cm(17.5))
    set_table_geometry(footer_table, [7200, 2780])
    left, right = footer_table.rows[0].cells
    set_font(left.paragraphs[0].add_run("TEG Holding · сформировано системой согласования"), size=7.5, color=MUTED)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_field(right.paragraphs[0])

    meta = doc.add_table(rows=2, cols=2)
    set_table_geometry(meta, [4990, 4990])
    set_table_borders(meta, color=LINE, size="5", inside=True)
    add_system_value(meta.cell(0, 0), "РЕГИСТРАЦИОННЫЙ НОМЕР", "{{DOCUMENT_NUMBER}}")
    add_system_value(meta.cell(0, 1), "ДАТА УТВЕРЖДЕНИЯ", "{{APPROVAL_DATE}}", WD_ALIGN_PARAGRAPH.RIGHT)
    add_system_value(meta.cell(1, 0), "ПРОЕКТ / ПОДРАЗДЕЛЕНИЕ", "{{PROJECT}}")
    add_system_value(meta.cell(1, 1), "ВСТУПАЕТ В СИЛУ", "{{EFFECTIVE_DATE}}", WD_ALIGN_PARAGRAPH.RIGHT)
    for cell in meta.rows[0].cells:
        set_cell_shading(cell, FILL)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(17)
    kicker.paragraph_format.space_after = Pt(4)
    set_font(kicker.add_run("ПРОЕКТ ДОКУМЕНТА"), size=8, bold=True, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(11)
    title.paragraph_format.keep_with_next = True
    set_font(title.add_run("{{DOCUMENT_TITLE}}"), size=16, bold=True, color=NAVY)

    note_table = doc.add_table(rows=1, cols=1)
    set_table_geometry(note_table, [CONTENT_WIDTH_DXA])
    set_table_borders(note_table, color="9EB4CF", size="7", inside=False)
    note = note_table.cell(0, 0)
    set_cell_shading(note, "EAF2FB")
    note_p = note.paragraphs[0]
    note_p.paragraph_format.space_after = Pt(0)
    set_font(note_p.add_run("ЗАПОЛНЯЕМАЯ ЧАСТЬ  "), size=8.5, bold=True, color=NAVY)
    set_font(
        note_p.add_run("Замените текст в синих полях. Не изменяйте реквизиты, название и системные обозначения выше."),
        size=8.5,
        color=BLUE,
    )

    action_word = "ПРИКАЗЫВАЮ" if kind == "order" else "ПОСТАНОВЛЯЕТ"
    action = doc.add_paragraph()
    action.alignment = WD_ALIGN_PARAGRAPH.CENTER
    action.paragraph_format.space_before = Pt(14)
    action.paragraph_format.space_after = Pt(6)
    set_font(action.add_run(action_word), size=11, bold=True, color=NAVY)

    add_section_heading(doc, "1", "ОСНОВНОЙ ТЕКСТ")
    add_fillable(doc, "Изложите решение руководства. Для нескольких требований используйте нумерованные пункты.", 6)

    add_section_heading(doc, "2", "ОТВЕТСТВЕННЫЕ ЛИЦА")
    add_fillable(doc, "Укажите ФИО, должности, задачи и сроки исполнения.", 3)

    add_section_heading(doc, "3", "ОСНОВАНИЕ")
    add_fillable(doc, "Укажите решение, документ или обстоятельство, послужившее основанием.", 2)

    add_section_heading(doc, "4", "ПРИЛОЖЕНИЯ")
    add_fillable(doc, "Перечислите приложения либо напишите «Нет».", 2)

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(12)
    closing.paragraph_format.space_after = Pt(0)
    set_font(closing.add_run("СЛУЖЕБНЫЕ РЕКВИЗИТЫ И ПОДПИСИ ДОБАВЛЯЮТСЯ СИСТЕМОЙ"), size=7.5, bold=True, color=MUTED)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output = OUTPUT_DIR / filename
    doc.save(output)
    print(output)


if __name__ == "__main__":
    build("order", "management_order_template.docx")
    build("resolution", "management_resolution_template.docx")
